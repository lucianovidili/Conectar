from django.views.generic import TemplateView, UpdateView
from django.views.generic.edit import CreateView, UpdateView, DeleteView
from django.shortcuts import render
from .forms import DonacionFormulario
from .models import Donacion
from django.urls import reverse_lazy
from django.core.paginator import Paginator
from django.http import JsonResponse
from rest_framework.pagination import PageNumberPagination
from rest_framework import generics
from .serializers import DonacionSerializer
from django.db.models import Q
from django.http import HttpResponse
import openpyxl
from django.utils.timezone import localtime
from openpyxl.styles import Font  # Importar Font para darle formato
from django.core.mail import send_mail
from django.conf import settings
from django.core.exceptions import ImproperlyConfigured
from smtplib import SMTPException


# INICIO


class inicio(TemplateView):
    template_name = "inicio.html"


# DONACIÓN


class crear_donacion(CreateView):
    model = Donacion
    form_class = DonacionFormulario
    success_url = reverse_lazy("ListaDonaciones")
    template_name = "crear_editar_donacion.html"


class editar_donacion(UpdateView):
    model = Donacion
    form_class = DonacionFormulario
    template_name = "crear_editar_donacion.html"
    success_url = reverse_lazy("ListaDonaciones")

    def form_valid(self, form):
        """Si el formulario es válido, guarda y redirige"""
        response = super().form_valid(form)
        return response

    def form_invalid(self, form):
        """Si el formulario no es válido, muestra errores y recarga la página"""
        print(form.errors)  # Para depuración en la terminal
        return self.render_to_response(self.get_context_data(form=form))


class eliminar_donacion(DeleteView):
    model = Donacion
    template_name = "eliminar_donacion.html"
    success_url = reverse_lazy("ListaDonaciones")


# class listar_jugadores_transferibles(LoginRequiredMixin, ListView):
# def lista_donaciones(request):
#     if (
#         False
#         # request.method == "GET"
#         # and request.headers.get("X-Requested-With") == "XMLHttpRequest"
#     ):
#         page_number = request.GET.get("page", 1)
#         donaciones = Donacion.objects.all().order_by("id")
#         paginator = Paginator(donaciones, 5)  # 5 jugadores por página

#         try:
#             page = paginator.page(page_number)
#         except:
#             return JsonResponse({"donaciones": [], "has_more": False})

#         donaciones_data = [
#             {
#                 "id": jugador.pk,
#                 "nombre": jugador.nombre,
#                 "apellido": jugador.apellido,
#                 "imagen_url": jugador.imagen.url if jugador.imagen else None,
#             }
#             for jugador in page
#         ]

#         # return JsonResponse({"jugadores": donaciones_data, "has_more": page.has_next()})
#         return render(request, "lista_donaciones.html", {"donaciones": donaciones})

#     donaciones = Donacion.objects.all().order_by("fecha_creacion")
#     return render(request, "lista_donaciones.html", {"donaciones": donaciones})


# def listar_donaciones(request):
#     donaciones = Donacion.objects.all().order_by("-fecha_creacion")
#     paginator = Paginator(donaciones, 10)  # 10 donaciones por página
#     page_number = request.GET.get("page", 1)
#     page_obj = paginator.get_page(page_number)

#     return render(request, "listar_donaciones.html", {"page_obj": page_obj})


def listar_donaciones(request):
    # Esta vista simplemente renderiza la plantilla HTML que contiene el frontend
    return render(request, "listar_donaciones.html")


def lista_donaciones(request):
    # Esta vista simplemente renderiza la plantilla HTML que contiene el frontend
    return render(request, "lista_donaciones.html")


def recuperar_contrasenia_solita(request, email):
    return render(request, "recuperar_contrasenia.html", {"email": email})


# def recuperar_contrasenia(request, email, password, titulo):
def recuperar_contrasenia(request, id):
    try:
        donacion = Donacion.objects.filter(id=id).first()
        email = donacion.email
        password = donacion.contrasenia
        titulo = donacion.titulo
        send_mail(
            subject="Recuperar contraseña",
            message=(
                "¡Hola! Te escribimos en nombre de SomosConectar. Te hemos enviado la contraseña para acceder a la donación titulada "
                + '"'
                + titulo
                + '"'
                + ". Tu contraseña es: "
                + password
                + "."
            ),
            from_email="postmaster@sandboxe3931eba496247ed946fb878f07ab861.mailgun.org",  # El correo de Mailgun
            recipient_list=[email],
            fail_silently=False,
        )
    except Exception as e:
        return render(
            request, "recuperar_contrasenia.html", {"email": email, "error": 1}
        )

    return render(request, "recuperar_contrasenia.html", {"email": email, "error": 0})


# Clase para manejar la paginación
class DonacionPagination(PageNumberPagination):
    page_size = 10  # Número de resultados por página
    page_size_query_param = (
        "page_size"  # Permite especificar el tamaño de la página con un parámetro
    )
    max_page_size = 100  # Tamaño máximo permitido por página


# Vista para listar las donaciones
class DonacionListView(generics.ListAPIView):
    queryset = Donacion.objects.all().order_by(
        "-fecha_creacion"
    )  # Ordenar por fecha de creación en forma descendente
    serializer_class = DonacionSerializer  # Usar el serializer que creamos
    pagination_class = DonacionPagination  # Usar la paginación que definimos


# def buscar_donaciones(request):
#     donaciones = Donacion.objects.all().values(
#         "id",
#         "titulo",
#         "descripcion",
#         "imagen",
#         "propietario",
#         "telefono",
#         "fecha_creacion",
#     )
#     return JsonResponse({"donaciones": list(donaciones)})


def buscar_donaciones(request):
    query = request.GET.get("q", "")
    page_number = request.GET.get("page", 1)

    donaciones = Donacion.objects.filter(
        Q(titulo__icontains=query)
        | Q(descripcion__icontains=query)
        | Q(propietario__icontains=query)
    ).order_by("-fecha_creacion")

    paginator = Paginator(donaciones, 10)  # 10 resultados por página
    page_obj = paginator.get_page(page_number)

    # Si la petición es AJAX, devolver JSON
    if request.headers.get("X-Requested-With") == "XMLHttpRequest":
        donaciones_data = list(
            page_obj.object_list.values(
                "titulo",
                "descripcion",
                "propietario",
                "id",
                "imagen",
                "telefono",
                "fecha_creacion",
            )
        )
        return JsonResponse(
            {"donaciones": donaciones_data, "has_next": page_obj.has_next()}
        )

    return render(
        request, "lista_donaciones.html", {"donaciones": page_obj, "query": query}
    )


def exportar_excel(request):
    # Crear archivo Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Donaciones"

    # Encabezados (excluyendo la 4ta y última columna)
    headers = ["FECHA CREACIÓN", "TÍTULO", "DESCRIPCIÓN", "PROPIETARIO", "TÉLEFONO"]
    ws.append(headers)

    # Obtener datos y convertir fecha a naive datetime
    donaciones = Donacion.objects.all().values_list(
        "fecha_creacion", "titulo", "descripcion", "propietario", "telefono"
    )

    for donacion in donaciones:
        fecha_sin_tz = localtime(donacion[0]).replace(tzinfo=None)  # Convertir a naive
        fecha_formateada = fecha_sin_tz.strftime(
            "%d/%m/%Y %H:%M:%S"
        )  # Formato "DD/MM/YYYY HH:MM:SS"
        fila = (fecha_formateada, *donacion[1:])  # Usar la fecha formateada en la tupla
        ws.append(fila)

    # Configurar la respuesta HTTP para la descarga
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="donaciones.xlsx"'
    wb.save(response)

    return response


# def exportar_excel(request):
# # Crear archivo Excel
# wb = openpyxl.Workbook()
# ws = wb.active
# ws.title = "Donaciones"

# # Encabezados (excluyendo la 4ta y última columna)
# headers = ["Fecha de Creación", "Título", "Descripción", "Propietario", "Teléfono"]

# # Establecer negrita en los encabezados
# font = Font(bold=True)  # Definir el estilo de negrita
# for col_num, header in enumerate(headers, 1):
#     cell = ws.cell(row=1, column=col_num)
#     cell.value = header
#     cell.font = font  # Aplicar el estilo de negrita

# # Obtener datos y convertir fecha a naive datetime
# donaciones = Donacion.objects.all().values_list(
#     "fecha_creacion", "titulo", "descripcion", "propietario", "telefono"
# )

# for donacion in donaciones:
#     fecha_sin_tz = localtime(donacion[0]).replace(tzinfo=None)  # Convertir a naive
#     fecha_formateada = fecha_sin_tz.strftime(
#         "%d/%m/%Y %H:%M:%S"
#     )  # Formatear como "DD/MM/YYYY HH:MM:SS"
#     fila = (
#         fecha_formateada,
#         *donacion[1:],
#     )  # Reemplazar la fecha formateada en la tupla
#     ws.append(fila)

# # Configurar la respuesta HTTP para la descarga
# response = HttpResponse(
#     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# )
# response["Content-Disposition"] = 'attachment; filename="donaciones.xlsx"'
# wb.save(response)

# return response
